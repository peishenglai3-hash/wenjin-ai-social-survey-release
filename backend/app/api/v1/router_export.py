"""Export endpoints — generate formatted documents from research content."""

import json
import logging
import os
import uuid
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, UploadFile
from fastapi.responses import Response, StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_ai_service, get_db, ensure_upload_dir
from app.api.deps_auth import get_current_profile
from app.models.survey import ExportDocument, SurveyProject, UserProfile
from app.services.ai import AIService
from app.services.export_service import generate_docx, generate_xlsx
from app.skills.mode_prompts import get_mode_modifier

router = APIRouter(tags=["export"])
logger = logging.getLogger(__name__)

# ── Reference file text extraction helpers ─────────────────────────

EXPORT_REFERENCE_ALLOWED = {".txt", ".md", ".docx", ".pdf"}
EXPORT_REFERENCE_MAX_MB = 10


def _extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract text from plain-text or markdown file, trying UTF-8 then GBK."""
    for enc in ("utf-8", "gbk", "gb2312", "latin-1"):
        try:
            return file_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a .docx file using python-docx."""
    from docx import Document
    doc = Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def _extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a .pdf file using pypdf."""
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def extract_reference_text(filename: str, file_bytes: bytes) -> dict:
    """Route to the correct extractor based on file extension. Returns {ok, text, error}."""
    suffix = Path(filename).suffix.lower()
    if suffix not in EXPORT_REFERENCE_ALLOWED:
        return {"ok": False, "text": "", "error": f"Unsupported format '{suffix}'. Supported: {', '.join(sorted(EXPORT_REFERENCE_ALLOWED))}"}

    size_mb = len(file_bytes) / (1024 * 1024)
    if size_mb > EXPORT_REFERENCE_MAX_MB:
        return {"ok": False, "text": "", "error": f"File too large ({size_mb:.1f}MB). Maximum: {EXPORT_REFERENCE_MAX_MB}MB"}

    try:
        if suffix in (".txt", ".md"):
            text = _extract_text_from_txt(file_bytes)
        elif suffix == ".docx":
            text = _extract_text_from_docx(file_bytes)
        elif suffix == ".pdf":
            text = _extract_text_from_pdf(file_bytes)
        else:
            return {"ok": False, "text": "", "error": f"Unsupported format: {suffix}"}

        if not text.strip():
            return {"ok": False, "text": "", "error": "No extractable text found in file (possibly scanned PDF or empty file)"}

        return {"ok": True, "text": text, "error": None}
    except Exception:
        logger.exception("Text extraction failed")
        return {"ok": False, "text": "", "error": "Text extraction failed"}


@router.post("/export/download-content")
async def export_download_content(payload: dict):
    """Generic download endpoint: convert any markdown content to .docx or .md file."""
    content = payload.get("content", "")
    file_type = payload.get("file_type", "docx")
    title = payload.get("title", "Export Document")

    if not content.strip():
        return {"error": "content is required"}

    safe_title = title.replace('"', '').replace('\n', ' ').strip() or "Export Document"

    if file_type == "docx":
        file_bytes = generate_docx(content, safe_title)
        encoded_filename = quote(f"{safe_title}.docx")
        return Response(
            content=file_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
                "Content-Length": str(len(file_bytes)),
            },
        )
    else:
        encoded = content.encode("utf-8")
        encoded_filename = quote(f"{safe_title}.md")
        return Response(
            content=encoded,
            media_type="text/markdown; charset=utf-8",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
                "Content-Length": str(len(encoded)),
            },
        )

EXPORT_PROMPTS = {
    # Academic Mode
    "paper_outline": (
        "Please organize the following research content into a standard academic paper outline, including:\n"
        "1. Title (Chinese and English)\n"
        "2. Abstract framework (Purpose-Method-Results-Conclusion)\n"
        "3. Introduction (research background, problem statement, significance)\n"
        "4. Literature review (theoretical framework and research hypotheses)\n"
        "5. Research methods (design, sample, tools, procedures)\n"
        "6. Research findings (organized by theme)\n"
        "7. Discussion (theoretical contributions, practical implications, limitations and future directions)\n"
        "8. Reference format suggestions"
    ),
    "report": (
        "Please organize the following research content into a social investigation report, including:\n"
        "1. Report title\n"
        "2. Investigation background and purpose\n"
        "3. Methodology description\n"
        "4. Core findings (use chart descriptions instead of actual data)\n"
        "5. Problem analysis and discussion\n"
        "6. Policy recommendations / Practical suggestions\n"
        "7. Conclusion"
    ),
    "proposal": (
        "Please organize the following research content into a research project proposal, including:\n"
        "1. Project name\n"
        "2. Rationale (research background, significance, current status)\n"
        "3. Research objectives and content\n"
        "4. Technical approach and methods\n"
        "5. Innovation points\n"
        "6. Expected outcomes\n"
        "7. Research timeline\n"
        "8. Budget framework"
    ),
    "reflection": (
        "Please organize the following research content into a practice reflection report, including:\n"
        "1. Practice process review\n"
        "2. Methodological reflection (was the method choice appropriate)\n"
        "3. Ethical reflection (informed consent, privacy protection, power relations)\n"
        "4. Personal growth and gains\n"
        "5. Reconsideration of social investigation work\n"
        "6. Future improvement directions"
    ),
    # Government Mode: Official Document Writing
    "official_document": (
        "Please write the following research content as a formal government document following the Party and Government Document Processing Regulations format:\n"
        "1. Title (issuing authority + subject + document type, e.g. 'Research Report on ...')\n"
        "2. Recipient (infer the appropriate recipient based on content)\n"
        "3. Introduction (background, purpose, basis, time, location, subjects, methods)\n"
        "4. Basic situation (summarize existing work foundation, main practices and results; use '[data to be filled]' as placeholder)\n"
        "5. Existing problems (list problems honestly, one paragraph per problem, analyze root causes)\n"
        "6. Policy recommendations (each recommendation should be specific, actionable, with clear responsible entity and implementation path)\n"
        "7. Conclusion (brief summary, express expectations for work advancement)\n"
        "8. Attachments (if supplementary materials are needed)\n\n"
        "Style requirements: Use formal government document language; solemn, accurate, concise; avoid empty rhetoric; emphasize problem orientation and operability."
    ),
    "policy_brief": (
        "Please write the following research content as a policy recommendation brief, format as follows:\n"
        "1. Brief title (concise, highlighting core recommendation)\n"
        "2. Background and problem (within 200 words, explain policy background and key issues)\n"
        "3. Current situation analysis (use data and cases to illustrate severity/urgency)\n"
        "4. Root cause analysis (analyze from institutional mechanisms, policy implementation, resource allocation dimensions)\n"
        "5. Policy recommendations (3-5 specific recommendations, each including: specific measures, implementing entity, expected effects, feasibility explanation)\n"
        "6. Risk assessment (potential resistance to implementation and coping strategies)\n"
        "Format: Concise language, clear priorities, easy for leadership approval; each recommendation no more than 300 words."
    ),
    # Youth Mode: Knowledge Summary
    "knowledge_summary": (
        "Please organize the following content into a knowledge summary for beginners to help understand core knowledge of social investigation and fieldwork:\n"
        "1. Core concepts (list 3-5 most important sociology/fieldwork concepts, each with accessible explanation and real-life examples)\n"
        "2. Knowledge context (show relationships between concepts in mind-map text form, why these concepts matter)\n"
        "3. Method tips (extract 3-5 practical fieldwork techniques, most common beginner mistakes and how to avoid them)\n"
        "4. Reflection questions (design 3-5 questions to guide readers in self-reflection, connecting knowledge with personal experience)\n"
        "5. Further reading (recommend 2-3 beginner-friendly resources with brief explanation of why worth reading)\n"
        "6. One-sentence summary (summarize the most important takeaway from this learning in one sentence)\n\n"
        "Style: Friendly, guiding, accessible. Explain every technical term. Use 'we' to build rapport. Encourage questioning and exploration."
    ),
    "learning_report": (
        "Please organize the following content into a youth learning exploration report:\n"
        "1. What I am curious about (personal motivation + social significance)\n"
        "2. What I observed and discovered (specific examples and observations)\n"
        "3. How I understand it through sociology (connect with classroom concepts, explain each concept)\n"
        "4. My reflections (my own assumptions, biases, and gains)\n"
        "5. What I still want to explore (unresolved questions and next learning directions)\n"
        "6. What I want to share (three key points if introducing to classmates)\n\n"
        "Style: First-person, friendly and natural, encouraging critical thinking."
    ),
}


# ── Reference file upload ──────────────────────────────────────────

@router.post("/export/upload-reference")
async def export_upload_reference(file: UploadFile = File(...)):
    """Upload a reference file (.txt/.md/.docx/.pdf) and extract its text.

    The extracted text can then be sent alongside the main content when
    generating export documents so the AI can reference it.
    """
    file_bytes = await file.read()
    result = extract_reference_text(file.filename or "upload", file_bytes)
    if result["ok"]:
        return {
            "status": "ok",
            "filename": file.filename,
            "text": result["text"],
            "text_length": len(result["text"]),
        }
    else:
        return {"status": "error", "filename": file.filename, "message": result["error"]}


@router.get("/export/formats")
async def export_formats():
    return {
        "formats": [
            # Academic
            {"id": "paper_outline", "label": "Paper Outline", "theme": "academic", "desc": "Standard academic paper chapter framework"},
            {"id": "report", "label": "Investigation Report", "theme": "academic", "desc": "Abstract-Method-Findings-Discussion-Recommendations"},
            {"id": "proposal", "label": "Project Proposal", "theme": "academic", "desc": "Research project/grant proposal writing assistance"},
            # Government
            {"id": "official_document", "label": "Official Document", "theme": "government", "desc": "Formal government document per Party and Government Document Processing Regulations"},
            {"id": "policy_brief", "label": "Policy Brief", "theme": "government", "desc": "Concise policy recommendations for leadership approval"},
            {"id": "gov_report", "label": "Government Research Report", "theme": "government", "desc": "Standard research report: situation-problems-recommendations"},
            # Youth
            {"id": "knowledge_summary", "label": "Knowledge Summary", "theme": "youth", "desc": "Core social investigation/fieldwork knowledge for beginners"},
            {"id": "learning_report", "label": "Learning Exploration Report", "theme": "youth", "desc": "First-person learning exploration and reflection report"},
            {"id": "reflection", "label": "Practice Reflection", "theme": "youth", "desc": "Social practice experience and reflection writing"},
        ],
    }


@router.post("/export/generate")
async def export_generate(
    payload: dict,
    profile: UserProfile = Depends(get_current_profile),
    ai: AIService = Depends(get_ai_service),
    db: AsyncSession = Depends(get_db),
):
    format_id = payload.get("format", "report")
    content = payload.get("content", "")
    additional_notes = payload.get("notes", "")
    project_id = payload.get("project_id", 0)
    theme = payload.get("theme", "academic")

    if not content.strip():
        return {"error": "content is required"}

    try:
        project_id = int(project_id or 0)
    except (TypeError, ValueError):
        project_id = 0
    if project_id:
        project_result = await db.execute(
            select(SurveyProject).where(
                SurveyProject.id == project_id,
                SurveyProject.profile_id == profile.id,
            )
        )
        if project_result.scalar_one_or_none() is None:
            return {"error": "Project not found"}

    export_prompt = EXPORT_PROMPTS.get(format_id, EXPORT_PROMPTS["report"])
    mode_export_hint = get_mode_modifier(theme, "export")
    mode_system_hint = get_mode_modifier(theme, "system")

    prompt = f"Research Content:\n{content}\n\nAdditional Notes: {additional_notes}\n\n{export_prompt}\n\n{mode_export_hint}"
    messages = [
        {"role": "system", "content": f"You are an academic writing assistant skilled at generating standardized documents from research content. {mode_system_hint} Output in Chinese, Markdown format."},
        {"role": "user", "content": prompt},
    ]
    result = await ai.chat(messages, temperature=0.5, max_tokens=4096, task="export")

    if project_id:
        doc = ExportDocument(project_id=project_id, format_type=format_id, content=result)
        db.add(doc)
        await db.commit()

    return {"format": format_id, "output": result}


@router.post("/export/generate/stream")
async def export_generate_stream(payload: dict, ai: AIService = Depends(get_ai_service)):
    """Stream document generation via SSE."""
    format_id = payload.get("format", "report")
    content = payload.get("content", "")
    additional_notes = payload.get("notes", "")
    theme = payload.get("theme", "academic")

    if not content.strip():
        return {"error": "content is required"}

    export_prompt = EXPORT_PROMPTS.get(format_id, EXPORT_PROMPTS["report"])
    mode_export_hint = get_mode_modifier(theme, "export")
    mode_system_hint = get_mode_modifier(theme, "system")

    prompt = f"Research Content:\n{content}\n\nAdditional Notes: {additional_notes}\n\n{export_prompt}\n\n{mode_export_hint}"
    messages = [
        {"role": "system", "content": f"You are an academic writing assistant skilled at generating standardized documents from research content. {mode_system_hint} Output in Chinese, Markdown format."},
        {"role": "user", "content": prompt},
    ]

    async def event_stream():
        async for token in ai.chat_stream(messages, temperature=0.5, max_tokens=4096, task="export"):
            yield f"data: {json.dumps({'token': token})}\n\n"
        yield "data: [DONE]\n\n"

    return StreamingResponse(event_stream(), media_type="text/event-stream")


@router.post("/export/generate-file")
async def export_generate_file(payload: dict, ai: AIService = Depends(get_ai_service)):
    """Generate and download a Word (.docx) or Excel (.xlsx) file."""
    format_id = payload.get("format", "report")
    file_type = payload.get("file_type", "docx")
    content = payload.get("content", "")

    if not content.strip():
        return {"error": "content is required"}

    export_prompt = EXPORT_PROMPTS.get(format_id, EXPORT_PROMPTS["report"])
    prompt = f"Research Content:\n{content}\n\n{export_prompt}"
    messages = [
        {"role": "system", "content": "You are an academic writing assistant skilled at generating standardized documents from research content. Output in Chinese, Markdown format."},
        {"role": "user", "content": prompt},
    ]
    ai_content = await ai.chat(messages, temperature=0.5, max_tokens=4096, task="export")

    if file_type == "xlsx":
        file_bytes = generate_xlsx(ai_content, title=format_id)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        encoded_filename = quote(f"{format_id}.xlsx")
    else:
        file_bytes = generate_docx(ai_content, title=format_id)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        encoded_filename = quote(f"{format_id}.docx")

    return Response(
        content=file_bytes,
        media_type=media_type,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
            "Content-Length": str(len(file_bytes)),
        },
    )


@router.get("/export/download/{export_id}")
async def export_download(
    export_id: int,
    format: str = "docx",
    profile: UserProfile = Depends(get_current_profile),
    db: AsyncSession = Depends(get_db),
):
    """Download an existing export record as a file."""
    result = await db.execute(
        select(ExportDocument)
        .join(SurveyProject, SurveyProject.id == ExportDocument.project_id)
        .where(
            ExportDocument.id == export_id,
            SurveyProject.profile_id == profile.id,
        )
    )
    doc = result.scalar_one_or_none()
    if not doc:
        return {"error": "Export not found"}

    if format == "xlsx":
        file_bytes = generate_xlsx(doc.content, title=doc.format_type)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        encoded_filename = quote(f"{doc.format_type}.xlsx")
    else:
        file_bytes = generate_docx(doc.content, title=doc.format_type)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        encoded_filename = quote(f"{doc.format_type}.docx")

    return Response(
        content=file_bytes,
        media_type=media_type,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
            "Content-Length": str(len(file_bytes)),
        },
    )
